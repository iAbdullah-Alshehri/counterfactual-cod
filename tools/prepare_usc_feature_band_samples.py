"""Build fresh, source-stratified USC12K train-only calibration and holdout manifests."""

from __future__ import annotations

import argparse
import csv
import json
import random
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from PIL import Image

from cod_recon.sinet import predict_images


def workbook_ids(path: Path) -> set[str]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    return {str(row[0]).strip() for row in sheet.iter_rows(min_row=2, values_only=True) if row[0]}


def all_manifest_ids(root: Path) -> set[str]:
    ids: set[str] = set()
    for manifest in root.glob("**/*manifest*.csv"):
        try:
            with manifest.open(newline="", encoding="utf-8-sig") as stream:
                for row in csv.DictReader(stream):
                    if row.get("image"):
                        ids.add(row.get("image_id") or Path(row["image"]).stem)
        except (OSError, csv.Error, KeyError):
            continue
    return ids


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--usc-root", type=Path, default=Path("data/raw/usc12k/VOC-USC12K"))
    parser.add_argument("--classes-dir", type=Path, default=Path("data/raw/usc12k_classes/Train_Class"))
    parser.add_argument("--cod-training-images", type=Path, default=Path("data/raw/cod_train/COD-TrainDataset/Imgs"))
    parser.add_argument("--checkpoint", type=Path, default=Path("data/downloads/SINet_V2_Net_epoch_best.pth"))
    parser.add_argument("--calibration-per-class", type=int, default=40)
    parser.add_argument("--holdout-per-class", type=int, default=15)
    parser.add_argument("--seed", type=int, default=20260928)
    parser.add_argument("--output-dir", type=Path, default=Path("data/experiments/usc12k_feature_band"))
    args = parser.parse_args()
    root, class_dir, out = args.usc_root.resolve(), args.classes_dir.resolve(), args.output_dir.resolve()
    if args.calibration_per_class < 30 or args.holdout_per_class < 1:
        parser.error("Use at least 30 calibration samples and at least one held-out sample per class")
    cal_manifest, hold_manifest = out / "calibration_manifest.csv", out / "heldout_manifest.csv"
    if cal_manifest.exists() or hold_manifest.exists():
        parser.error(f"Refusing to overwrite existing sample manifests in {out}")

    train_ids = set((root / "ImageSets/Segmentation/train.txt").read_text(encoding="utf-8").split())
    official_test_ids = set((root / "ImageSets/Segmentation/val.txt").read_text(encoding="utf-8").split())
    if train_ids & official_test_ids:
        parser.error("USC training and official test ID lists overlap")
    scenes = {name: workbook_ids(class_dir / f"Scene{name}2.xlsx") for name in "ABC"}
    if any(not values <= train_ids for values in scenes.values()):
        parser.error("A USC scene workbook contains IDs outside train.txt")
    pools = {"CO": scenes["C"], "BG": train_ids - scenes["A"] - scenes["B"] - scenes["C"], "NOCOD": scenes["A"]}
    excluded = all_manifest_ids(Path("data/experiments"))
    cod_stems = {p.stem for p in args.cod_training_images.rglob("*") if p.is_file()}
    image_root, label_root = root / "JPEGImages", root / "SegmentationClass"
    eligible = {label: {i for i in values - excluded - cod_stems - official_test_ids
                        if (image_root / f"{i}.jpg").is_file()}
                for label, values in pools.items()}

    # Preserve the observed USC source mix rather than selecting only one source
    # per class: tiny CO sources get representation, and the two-source BG/NOCOD
    # pools are sampled approximately evenly.
    calibration_plan = {"CO": {"CS-": 35, "AWA": 3, "LUSI_": 2},
                        "BG": {"COD10K-NonCAM-": 20, "background-": 20},
                        "NOCOD": {"DUTS-": 20, "HKU-": 20}}
    holdout_plan = {"CO": {"CS-": 13, "AWA": 1, "LUSI_": 1},
                    "BG": {"COD10K-NonCAM-": 8, "background-": 7},
                    "NOCOD": {"DUTS-": 8, "HKU-": 7}}
    if args.calibration_per_class != 40 or args.holdout_per_class != 15:
        parser.error("The documented source-stratified plan currently requires 40 calibration and 15 holdout per class")
    rng = random.Random(args.seed)
    selected: dict[str, dict[str, list[str]]] = {"calibration": {}, "heldout": {}}
    used = set(excluded)
    for split_name, plan in (("calibration", calibration_plan), ("heldout", holdout_plan)):
        for label, source_counts in plan.items():
            selected[split_name][label] = []
            for prefix, count in source_counts.items():
                candidates = sorted(i for i in eligible[label] if i.startswith(prefix) and i not in used)
                if len(candidates) < count:
                    parser.error(f"Insufficient {label}/{prefix} candidates for {split_name}: {len(candidates)} < {count}")
                rng.shuffle(candidates)
                chosen = candidates[:count]
                selected[split_name][label].extend(chosen)
                used.update(chosen)
    if set(sum((sum(labels.values(), []) for labels in selected.values()), [])) & excluded:
        parser.error("New selections overlap an existing experiment manifest")
    if set(selected["calibration"]["CO"] + selected["calibration"]["BG"] + selected["calibration"]["NOCOD"]) & set(sum((v for v in selected["heldout"].values()), [])):
        parser.error("Calibration and heldout samples overlap")

    out.mkdir(parents=True, exist_ok=True)
    all_images = [image_root / f"{image_id}.jpg"
                  for split in ("calibration", "heldout")
                  for label in ("CO", "BG", "NOCOD") for image_id in selected[split][label]]
    index = predict_images(all_images, image_root, out / "candidate_masks", args.checkpoint.resolve())
    with index.open(newline="", encoding="utf-8") as stream:
        predictions = {Path(r["image"]).stem: r["mask"] for r in csv.DictReader(stream)}

    def source_name(image_id: str) -> str:
        for prefix, name in (("CS-", "CS"), ("AWA", "AWA"), ("LUSI_", "LUSI"),
                             ("COD10K-NonCAM-", "COD10K"), ("background-", "BACKGROUND"),
                             ("DUTS-", "DUTS"), ("HKU-", "HKU")):
            if image_id.startswith(prefix):
                return name
        return "unknown"

    def write_manifest(split_name: str) -> None:
        filename = cal_manifest if split_name == "calibration" else hold_manifest
        gt_dir = out / "ground_truth" / split_name
        gt_dir.mkdir(parents=True, exist_ok=True)
        with filename.open("w", newline="", encoding="utf-8") as stream:
            fields = ("image", "label", "mask", "ground_truth", "split", "image_id", "source_family")
            writer = csv.DictWriter(stream, fieldnames=fields)
            writer.writeheader()
            for label in ("CO", "BG", "NOCOD"):
                for image_id in selected[split_name][label]:
                    gt = ""
                    if label == "CO":
                        with Image.open(label_root / f"{image_id}.png") as indexed:
                            palette_mask = np.asarray(indexed.convert("P"))
                            gt_path = gt_dir / f"{image_id}.png"
                            Image.fromarray(np.where(palette_mask == 2, 255, 0).astype(np.uint8), mode="L").save(gt_path)
                        gt = str(gt_path)
                    writer.writerow({"image": str((image_root / f"{image_id}.jpg").resolve()),
                                     "label": label, "mask": predictions[image_id], "ground_truth": gt,
                                     "split": f"usc12k_train_feature_{split_name}", "image_id": image_id,
                                     "source_family": source_name(image_id)})
        print(f"Wrote {filename}")

    write_manifest("calibration")
    write_manifest("heldout")
    metadata = {"seed": args.seed, "calibration_n_per_class": args.calibration_per_class,
                "heldout_n_per_class": args.holdout_per_class, "prior_manifest_ids_excluded": len(excluded),
                "sinet_training_stems_excluded": len(cod_stems), "official_test_images_used": False,
                "calibration_sources": {k: {s: len([i for i in vals if source_name(i) == s])
                                             for s in sorted({source_name(i) for i in vals})}
                                         for k, vals in selected["calibration"].items()},
                "heldout_sources": {k: {s: len([i for i in vals if source_name(i) == s])
                                        for s in sorted({source_name(i) for i in vals})}
                                    for k, vals in selected["heldout"].items()},
                "calibration_ids": selected["calibration"], "heldout_ids": selected["heldout"]}
    (out / "sample_selection.json").write_text(json.dumps(metadata, indent=2) + "\n", encoding="utf-8")
    print(f"Excluded IDs from existing manifests: {len(excluded)}")
    print(f"Excluded exact-stem SINet-V2 training IDs: {len(cod_stems)}")
    print(f"Calibration 40/class; holdout 15/class; official USC test overlap=0; sample sets mutually disjoint")
    print(f"Calibration sources: {metadata['calibration_sources']}")
    print(f"Heldout sources: {metadata['heldout_sources']}")


if __name__ == "__main__":
    main()
