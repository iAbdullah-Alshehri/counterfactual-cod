"""Prepare fresh, disjoint COD10K development and USC12K sanity samples."""

from __future__ import annotations

import argparse
import csv
import random
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from PIL import Image

from cod_recon.sinet import predict_images


def workbook_ids(path: Path) -> set[str]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    return {str(row[0]).strip() for row in sheet.iter_rows(min_row=2, values_only=True) if row[0]}


def manifest_ids(path: Path) -> set[str]:
    if not path.is_file():
        return set()
    with path.open(newline="", encoding="utf-8-sig") as stream:
        return {row.get("image_id") or Path(row["image"]).stem for row in csv.DictReader(stream)}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--usc-root", type=Path, default=Path("data/raw/usc12k/VOC-USC12K"))
    parser.add_argument("--usc-classes", type=Path, default=Path("data/raw/usc12k_classes/Train_Class"))
    parser.add_argument("--cod-test-root", type=Path, default=Path("data/raw/cod_test/COD-TestDataset/COD10K"))
    parser.add_argument("--cod-training-images", type=Path, default=Path("data/raw/cod_train/COD-TrainDataset/Imgs"))
    parser.add_argument("--checkpoint", type=Path, default=Path("data/downloads/SINet_V2_Net_epoch_best.pth"))
    parser.add_argument("--dev-per-class", type=int, default=25)
    parser.add_argument("--sanity-per-class", type=int, default=15)
    parser.add_argument("--seed", type=int, default=20260927)
    parser.add_argument("--output-dir", type=Path, default=Path("data/experiments/cod10k_feature_only"))
    args = parser.parse_args()
    usc, cls = args.usc_root.resolve(), args.usc_classes.resolve()
    train_ids = set((usc / "ImageSets/Segmentation/train.txt").read_text(encoding="utf-8").split())
    official_test_ids = set((usc / "ImageSets/Segmentation/val.txt").read_text(encoding="utf-8").split())
    if train_ids & official_test_ids:
        parser.error("USC train and official test IDs overlap")
    scenes = {name: workbook_ids(cls / f"Scene{name}2.xlsx") for name in "ABC"}
    if any(not ids <= train_ids for ids in scenes.values()):
        parser.error("A scene workbook contains IDs outside USC12K train.txt")
    pools = {"NOCOD": scenes["A"], "CO": scenes["C"],
             "BG": train_ids - scenes["A"] - scenes["B"] - scenes["C"]}
    cod_stems = {p.stem for p in args.cod_training_images.rglob("*") if p.is_file()}
    prior = (manifest_ids(Path("data/experiments/cod10k_matched/matched_manifest.csv")) |
             manifest_ids(Path("data/experiments/usc12k/validation/validation_manifest.csv")))
    # This extra glob catches earlier exploratory manifests while keeping raw data untouched.
    for path in Path("data/experiments").glob("**/*manifest*.csv"):
        if path.name not in {"matched_manifest.csv", "validation_manifest.csv"} and path.parent != args.output_dir:
            prior |= manifest_ids(path)
    rng = random.Random(args.seed)

    codroot = args.cod_test_root.resolve()
    positives = [p for p in (codroot / "Imgs").glob("COD10K-CAM-*.jpg")
                 if (codroot / "GT" / f"{p.stem}.png").is_file()]
    positives = [p for p in positives if p.stem not in cod_stems and p.stem not in prior]
    negative_ids = {i for i in pools["BG"] if i.startswith("COD10K-NonCAM-")
                    and i not in cod_stems and i not in prior
                    and (usc / "JPEGImages" / f"{i}.jpg").is_file()}
    if min(len(positives), len(negative_ids)) < args.dev_per_class:
        parser.error(f"Insufficient fresh COD10K data: CO={len(positives)}, NonCAM={len(negative_ids)}")
    rng.shuffle(positives)
    fresh_pos = positives[:args.dev_per_class]
    fresh_neg = sorted(negative_ids)
    rng.shuffle(fresh_neg)
    fresh_neg = fresh_neg[:args.dev_per_class]

    # Use one source family per class for transparent source reporting. Keep all
    # selections disjoint from prior manifests and the just-selected COD pool.
    used = prior | {p.stem for p in fresh_pos} | set(fresh_neg)
    usc_sources = {
        "CO": [i for i in pools["CO"] if i.startswith("CS-")],
        "BG": [i for i in pools["BG"] if i.startswith("background-")],
        "NOCOD": [i for i in pools["NOCOD"] if i.startswith("HKU-")],
    }
    sanity: dict[str, list[str]] = {}
    for label in ("CO", "BG", "NOCOD"):
        candidates = sorted(i for i in usc_sources[label] if i not in used and i not in cod_stems
                            and (usc / "JPEGImages" / f"{i}.jpg").is_file())
        if len(candidates) < args.sanity_per_class:
            parser.error(f"Insufficient disjoint USC {label}={len(candidates)}")
        rng.shuffle(candidates)
        sanity[label] = candidates[:args.sanity_per_class]
        used.update(sanity[label])
    sanity_ids = set(sum(sanity.values(), []))
    if sanity_ids & official_test_ids:
        parser.error("USC sanity set overlaps official test IDs")

    out = args.output_dir.resolve()
    out.mkdir(parents=True, exist_ok=True)
    image_root, gt_root = usc / "JPEGImages", usc / "SegmentationClass"

    def write_manifest(filename: str, rows: list[dict], mask_subdir: str) -> None:
        predicted = {}
        roots = {"cod": codroot / "Imgs", "usc": image_root}
        for source, root in roots.items():
            images = [r["path"] for r in rows if r["root"] == source]
            if not images:
                continue
            index = predict_images(images, root, out / "candidate_masks" / mask_subdir / source,
                                  args.checkpoint.resolve())
            with index.open(encoding="utf-8") as stream:
                for row in csv.DictReader(stream):
                    predicted[Path(row["image"]).stem] = row["mask"]
        target_dir = out / "ground_truth" / mask_subdir
        target_dir.mkdir(parents=True, exist_ok=True)
        path = out / filename
        fields = ("image", "label", "mask", "ground_truth", "split", "image_id", "source_family", "source_subset")
        with path.open("w", newline="", encoding="utf-8") as stream:
            writer = csv.DictWriter(stream, fieldnames=fields)
            writer.writeheader()
            for row in rows:
                image = row["path"]
                label, image_id = row["label"], image.stem
                gt = ""
                if label == "CO":
                    if row["root"] == "cod":
                        gt = str((codroot / "GT" / f"{image_id}.png").resolve())
                    else:
                        gt_path = target_dir / f"{image_id}.png"
                        with Image.open(gt_root / f"{image_id}.png") as indexed:
                            palette_mask = np.asarray(indexed.convert("P"))
                            Image.fromarray(np.where(palette_mask == 2, 255, 0).astype(np.uint8), mode="L").save(gt_path)
                        gt = str(gt_path)
                writer.writerow({"image": str(image), "label": label, "mask": predicted[image_id],
                                 "ground_truth": gt, "split": row["split"], "image_id": image_id,
                                 "source_family": row["family"], "source_subset": row["subset"]})
        print(f"Wrote {path}")

    dev_rows = ([{"path": p, "root": "cod", "label": "CO", "split": "fresh_cod10k_feature_development", "family": "COD10K", "subset": "official_COD10K_test_CAM"} for p in fresh_pos] +
                [{"path": image_root / f"{i}.jpg", "root": "usc", "label": "BG", "split": "fresh_cod10k_feature_development", "family": "COD10K", "subset": "USC12K_train_COD10K_NonCAM"} for i in fresh_neg])
    write_manifest("dev_manifest.csv", dev_rows, "cod10k")
    source_names = {"CO": "USC12K_train_CS", "BG": "USC12K_train_BACKGROUND", "NOCOD": "USC12K_train_HKU"}
    sanity_rows = [{"path": image_root / f"{i}.jpg", "label": label,
                    "root": "usc", "split": "heldout_usc12k_train_sanity", "family": label,
                    "subset": source_names[label]}
                   for label in ("CO", "BG", "NOCOD") for i in sanity[label]]
    write_manifest("usc_sanity_manifest.csv", sanity_rows, "usc_sanity")
    report = {"seed": args.seed, "dev_per_class": args.dev_per_class,
              "sanity_per_class": args.sanity_per_class, "prior_ids_excluded": len(prior),
              "sinet_training_stems_excluded": len(cod_stems),
              "dev_ids": [p.stem for p in fresh_pos] + fresh_neg,
              "sanity_ids": {label: ids for label, ids in sanity.items()},
              "official_test_id_overlap": 0}
    (out / "sample_selection.json").write_text(__import__("json").dumps(report, indent=2) + "\n", encoding="utf-8")
    print(f"Fresh COD dev sample: CO={len(fresh_pos)}, NonCAM={len(fresh_neg)}; excluded prior IDs={len(prior)}")
    print("USC sanity sample: " + ", ".join(f"{k}={len(v)}" for k, v in sanity.items()))
    print(f"Overlap with SINet training archive excluded; official USC test overlap=0")


if __name__ == "__main__":
    main()
