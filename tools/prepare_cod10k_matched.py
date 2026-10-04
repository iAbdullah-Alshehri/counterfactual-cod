"""Prepare COD10K camouflaged vs same-source NonCAM samples for a shortcut check."""

from __future__ import annotations

import argparse
import csv
import random
from pathlib import Path

from openpyxl import load_workbook

from cod_recon.sinet import predict_images


def _class_ids(path: Path) -> set[str]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    return {str(row[0]).strip() for row in sheet.iter_rows(min_row=2, values_only=True) if row[0]}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--usc-root", type=Path, default=Path("data/raw/usc12k/VOC-USC12K"))
    parser.add_argument("--usc-classes", type=Path, default=Path("data/raw/usc12k_classes/Train_Class"))
    parser.add_argument("--cod10k-test-root", type=Path, default=Path("data/raw/cod_test/COD-TestDataset/COD10K"))
    parser.add_argument("--cod-training-images", type=Path, default=Path("data/raw/cod_train/COD-TrainDataset/Imgs"))
    parser.add_argument("--checkpoint", type=Path, default=Path("data/downloads/SINet_V2_Net_epoch_best.pth"))
    parser.add_argument("--sample-per-class", type=int, default=25)
    parser.add_argument("--seed", type=int, default=20260925)
    parser.add_argument("--output-dir", type=Path, default=Path("data/experiments/cod10k_matched"))
    args = parser.parse_args()
    if args.sample_per_class < 1:
        parser.error("--sample-per-class must be positive")

    usc, classes = args.usc_root.resolve(), args.usc_classes.resolve()
    train_ids = set((usc / "ImageSets/Segmentation/train.txt").read_text(encoding="utf-8").split())
    scene_ids = [_class_ids(classes / name) for name in ("SceneA2.xlsx", "SceneB2.xlsx", "SceneC2.xlsx")]
    negative_pool = {image_id for image_id in train_ids - set.union(*scene_ids)
                     if image_id.startswith("COD10K-NonCAM-")}
    negative_root = usc / "JPEGImages"
    negative_pool = {image_id for image_id in negative_pool if (negative_root / f"{image_id}.jpg").is_file()}

    cod_test = args.cod10k_test_root.resolve()
    positive_root, positive_gt = cod_test / "Imgs", cod_test / "GT"
    positive_paths = [path for path in positive_root.glob("COD10K-CAM-*.jpg")
                      if (positive_gt / f"{path.stem}.png").is_file()]
    positive_paths.sort()
    positive_stems = {path.stem for path in positive_paths}
    training_stems = {path.stem for path in args.cod_training_images.rglob("*") if path.is_file()}
    excluded = positive_stems & training_stems
    positive_paths = [path for path in positive_paths if path.stem not in training_stems]
    negative_pool -= training_stems
    if len(positive_paths) < args.sample_per_class or len(negative_pool) < args.sample_per_class:
        parser.error(f"Insufficient non-overlapping data: CO={len(positive_paths)}, BG={len(negative_pool)}")
    if {path.stem for path in positive_paths} & negative_pool:
        parser.error("Positive and negative image IDs overlap")

    rng = random.Random(args.seed)
    rng.shuffle(positive_paths)
    negative_ids = sorted(negative_pool)
    rng.shuffle(negative_ids)
    positive_paths = positive_paths[:args.sample_per_class]
    negative_ids = negative_ids[:args.sample_per_class]
    out = args.output_dir.resolve()
    out.mkdir(parents=True, exist_ok=True)
    positive_index = predict_images(positive_paths, positive_root, out / "candidate_masks/COD10K_CAM",
                                   args.checkpoint.resolve())
    negative_paths = [negative_root / f"{image_id}.jpg" for image_id in negative_ids]
    negative_index = predict_images(negative_paths, negative_root, out / "candidate_masks/COD10K_NonCAM",
                                   args.checkpoint.resolve())
    prediction = {}
    for index in (positive_index, negative_index):
        for row in csv.DictReader(index.open(encoding="utf-8")):
            prediction[Path(row["image"]).resolve()] = Path(row["mask"]).resolve()

    manifest = out / "matched_manifest.csv"
    with manifest.open("w", newline="", encoding="utf-8") as stream:
        fields = ("image", "label", "mask", "ground_truth", "split", "image_id", "source_family", "source_subset")
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        for path in positive_paths:
            writer.writerow({"image": str(path.resolve()), "label": "CO",
                             "mask": str(prediction[path.resolve()]),
                             "ground_truth": str((positive_gt / f"{path.stem}.png").resolve()),
                             "split": "cod10k_source_matched_validation", "image_id": path.stem,
                             "source_family": "COD10K", "source_subset": "official_COD10K_test_CAM"})
        for path in negative_paths:
            writer.writerow({"image": str(path.resolve()), "label": "BG",
                             "mask": str(prediction[path.resolve()]), "ground_truth": "",
                             "split": "cod10k_source_matched_validation", "image_id": path.stem,
                             "source_family": "COD10K", "source_subset": "USC12K_train_COD10K_NonCAM"})
    print(f"COD10K CAM candidates: {len(positive_stems)}; exact-stem overlap with SINet train excluded: {len(excluded)}")
    print(f"Eligible COD10K NonCAM in USC12K train: {len(negative_pool)}")
    print(f"Wrote {len(positive_paths)} CO and {len(negative_paths)} matched-source BG rows: {manifest}")


if __name__ == "__main__":
    main()
